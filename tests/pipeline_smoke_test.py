#!/usr/bin/env python3
"""Smoke test for suggest_column_merges.py + merge_to_master.py (hardened).

Builds synthetic fixtures reproducing every failure pattern found in the
audit, runs both scripts, and asserts:

  suggest_column_merges.py
    1. exit 0 with control characters in headers/values (no late crash)
    2. no live formulas in the suggestions workbook (injection neutralized)
    3. cp1252 CSV profiled (encoding fallback) — its values appear as samples
    4. unreadable file recorded in the Skipped_Files sheet

  merge_to_master.py
    5. exit 0 on the same dirty inputs
    6. exact master row count; same-row merge rule (identical / different /
       blank-first) verified
    7. unmapped column data never enters the master (allowlist)
    8. cp1252 rows ARE ingested, with correct accents
    9. no live formulas anywhere, including Unmapped_Headers (=SUM header
       stored as text)
   10. every failure and preview-part skip appears in the Ingest_Log
   11. *_partNNN.xlsx preview files skipped by default, while ordinary
       preview-format files are still ingested
   12. Excel row-limit guard: an oversized master splits into Master,
       Master_2, ... sheets (tested with a lowered limit)

Usage: python tests/pipeline_smoke_test.py
"""

import contextlib
import io
import os
import shutil
import subprocess
import sys
import tempfile

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(HERE, ".."))
SUGGEST = os.path.join(REPO, "suggest_column_merges.py")
MERGE = os.path.join(REPO, "merge_to_master.py")

FAILURES = []


def check(label, condition, detail=""):
    status = "PASS" if condition else "FAIL"
    print(f"  [{status}] {label}" + (f" — {detail}" if detail and not condition else ""))
    if not condition:
        FAILURES.append(label)


def build_fixtures(base):
    os.makedirs(base, exist_ok=True)

    # f1: clean file; "Secret Notes" stays unmapped (allowlist test)
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Leads"
    ws.append(["Company Name", "Owner Name", "Mobile", "City", "Secret Notes"])
    for i in range(1, 6):
        ws.append([f"DummyCo {i} Pvt Ltd", f"Owner {i}", f"90000000{i:02d}", "Mumbai", "TOPSECRET123"])
    wb.save(os.path.join(base, "f1_leads.xlsx"))

    # f3: two phone columns in one sheet (same-row merge rule)
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Phones"
    ws.append(["Company", "Phone 1", "Phone 2", "City"])
    ws.append(["Acme Pvt Ltd", "9111100001", "9111100001", "Mumbai"])
    ws.append(["Beta Traders", "9222200002", "02233445566", "Pune"])
    ws.append(["Gamma Corp", "", "9333300003", "Delhi"])
    wb.save(os.path.join(base, "f3_phones.xlsx"))

    # f4 + pv_part001: identical preview-format content; only the part-named
    # one must be skipped by merge_to_master's default
    for name in ("f4_preview.xlsx", "pv_part001.xlsx"):
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = "prev"
        ws.append(["Field", "Value"]); ws.append(["file_path", "orig/f.xlsx"])
        ws.append(["file_name", "f.xlsx"]); ws.append(["sheet_name", "Orig"])
        ws.append(["header_row_no", 1]); ws.append(["header_detected", "yes"])
        ws.append(["total_columns", 2]); ws.append(["sample_rows", 2])
        ws.append(["Company Name", "Mobile"])
        ws.append(["Preview Co One", "9555500001"])
        ws.append(["Preview Co Two", "9555500002"])
        wb.save(os.path.join(base, name))

    # f5: headerless CSV (misdetection -> skipped as unmapped, logged)
    with open(os.path.join(base, "f5_headerless.csv"), "w", encoding="utf-8") as f:
        f.write("Delta Corp,9444400004,Nagpur\nEpsilon Ltd,9444400005,Akola\n")

    # f6: cp1252 CSV — must be ingested via encoding fallback
    with open(os.path.join(base, "f6_cp1252.csv"), "wb") as f:
        f.write("Company Name,City\nCaf\xe9 Exports,Delhi\nPlain Co,Pune\n".encode("cp1252"))

    # f8: control characters in a header and a value
    with open(os.path.join(base, "f8_ctrl.csv"), "wb") as f:
        f.write(b"bad\x01col,City\nhas\x01ctrl value,Delhi\nclean value,Pune\n")

    # f9: formula payloads in a value AND in an (unmapped) header
    with open(os.path.join(base, "f9_inject.csv"), "w", encoding="utf-8") as f:
        f.write('Company Name,Remark Col,"=SUM(A1:A99)"\n')
        f.write('"Inject Co","=HYPERLINK(""http://evil.example/z"",""open me"")","x"\n')

    # f11: corrupt xlsx (failure accounting)
    with open(os.path.join(base, "f11_corrupt.xlsx"), "wb") as f:
        f.write(b"this is not a zip archive")


def write_mapping(path):
    with open(path, "w", encoding="utf-8") as f:
        f.write("raw_column_name,suggested_canonical\n")
        for raw, canon in [
            ("Company Name", "company_name"), ("Owner Name", "person_name"),
            ("Mobile", "phone"), ("City", "city"),
            ("Company", "company_name"), ("Phone 1", "phone"), ("Phone 2", "phone"),
            ("Remark Col", "other_text"),
        ]:
            f.write(f"{raw},{canon}\n")


def workbook_cells(path):
    wb = openpyxl.load_workbook(path)
    cells = []
    for name in wb.sheetnames:
        for row in wb[name].iter_rows():
            for c in row:
                if c.value is not None:
                    cells.append((name, c))
    return wb, cells


def main():
    workdir = tempfile.mkdtemp(prefix="pipeline_smoke_")
    fixtures = os.path.join(workdir, "fixtures")
    build_fixtures(fixtures)
    mapping = os.path.join(workdir, "mapping.csv")
    write_mapping(mapping)
    suggestions = os.path.join(workdir, "suggestions.xlsx")
    master = os.path.join(workdir, "master.xlsx")

    # ── suggest_column_merges ────────────────────────────────────────────────
    print("Running suggest_column_merges...")
    p1 = subprocess.run([sys.executable, SUGGEST, fixtures, "--output", suggestions],
                        capture_output=True, text=True)
    print("Checks (suggest):")
    check("exit 0 despite control chars", p1.returncode == 0,
          f"exit={p1.returncode}\n{p1.stdout[-800:]}\n{p1.stderr[-800:]}")
    if p1.returncode != 0:
        finish(workdir)

    wb1, cells1 = workbook_cells(suggestions)
    text1 = " | ".join(str(c.value) for _, c in cells1)
    check("no live formulas in suggestions",
          all(c.data_type != "f" for _, c in cells1))
    check("HYPERLINK payload present only as text", "HYPERLINK" in text1)
    check("control char stripped, value kept", "hasctrl value" in text1)
    check("cp1252 file profiled (accents correct)", "Café Exports" in text1)
    check("Skipped_Files sheet lists corrupt file",
          "Skipped_Files" in wb1.sheetnames
          and any("f11_corrupt" in str(r[0].value) for r in wb1["Skipped_Files"].iter_rows(min_row=2)))

    # ── merge_to_master ──────────────────────────────────────────────────────
    print("\nRunning merge_to_master...")
    p2 = subprocess.run(
        [sys.executable, MERGE, "--input", fixtures, "--mapping", mapping,
         "--output", master, "--include-source"],
        capture_output=True, text=True)
    print("Checks (merge):")
    check("exit 0 despite dirty inputs", p2.returncode == 0,
          f"exit={p2.returncode}\n{p2.stdout[-800:]}\n{p2.stderr[-800:]}")
    if p2.returncode != 0:
        finish(workdir)

    wb2, cells2 = workbook_cells(master)
    text2 = " | ".join(str(c.value) for _, c in cells2)
    m = wb2["Master"]
    hdr = [c.value for c in m[1]]
    rows = [dict(zip(hdr, r)) for r in m.iter_rows(min_row=2, values_only=True)]

    # expected: f1=5, f3=3, f4=2, f6=2, f8=2 (City col mapped), f9=1;
    # f5 unmapped-skip, pv_part skipped, f11 error
    check("master row count exact (15)", len(rows) == 15, f"got {len(rows)}")
    check("no live formulas anywhere in master workbook",
          all(c.data_type != "f" for _, c in cells2))
    check("unmapped column data excluded (allowlist)", "TOPSECRET123" not in text2)
    check("cp1252 rows ingested with correct accents",
          any(r.get("company_name") == "Café Exports" for r in rows))
    check("preview-format file ingested",
          sum(1 for r in rows if r.get("source_file") == "f4_preview.xlsx") == 2)
    check("preview part file NOT ingested",
          not any(r.get("source_file") == "pv_part001.xlsx" for r in rows))

    phones = {r.get("company_name"): r.get("phone")
              for r in rows if r.get("source_file") == "f3_phones.xlsx"}
    check("merge rule: identical values collapse", phones.get("Acme Pvt Ltd") == "9111100001")
    check("merge rule: different values joined",
          phones.get("Beta Traders") == "9222200002 | 02233445566")
    check("merge rule: blank-first falls through", phones.get("Gamma Corp") == "9333300003")

    unm = [str(r[0].value) for r in wb2["Unmapped_Headers"].iter_rows(min_row=2)]
    check("=SUM header in Unmapped_Headers escaped to text",
          any("SUM(A1:A99)" in u for u in unm)
          and all(c.data_type != "f" for c in wb2["Unmapped_Headers"]["A"]))
    check("ctrl-char header sanitized in Unmapped_Headers",
          any(u == "badcol" for u in unm), str(unm))

    lh = [c.value for c in wb2["Ingest_Log"][1]]
    log = [dict(zip(lh, r)) for r in wb2["Ingest_Log"].iter_rows(min_row=2, values_only=True)]
    check("Ingest_Log: corrupt file recorded as error",
          any("f11_corrupt" in str(r["file_name"]) and r["status"] == "error" for r in log))
    check("Ingest_Log: preview part recorded as skipped",
          any("pv_part001" in str(r["file_name"]) and r["status"] == "skipped" for r in log))
    check("Ingest_Log: headerless file recorded as skipped",
          any("f5_headerless" in str(r["file_name"]) and r["status"] == "skipped" for r in log))
    check("Ingest_Log: cp1252 encoding noted",
          any("f6" in str(r["file_name"]) and "cp1252" in str(r["detail"]) for r in log))

    # ── Excel row-limit chunking (with a lowered limit) ──────────────────────
    print("\nChecking Master sheet chunking...")
    sys.path.insert(0, REPO)
    import merge_to_master as mtm
    original_limit = mtm.EXCEL_MAX_ROWS
    mtm.EXCEL_MAX_ROWS = 6  # 5 data rows per sheet
    chunk_out = os.path.join(workdir, "chunked.xlsx")
    try:
        with contextlib.redirect_stdout(io.StringIO()):
            mtm.run(inputs=[fixtures], mapping_path=mapping, output_path=chunk_out,
                    overwrite=True)
    finally:
        mtm.EXCEL_MAX_ROWS = original_limit
    wb3 = openpyxl.load_workbook(chunk_out)
    masters = [s for s in wb3.sheetnames if s.startswith("Master")]
    data_rows = sum(wb3[s].max_row - 1 for s in masters)
    check("oversized master split across sheets", masters == ["Master", "Master_2", "Master_3"],
          str(wb3.sheetnames))
    check("no rows lost in the split", data_rows == 15, f"got {data_rows}")

    finish(workdir)


def finish(workdir):
    shutil.rmtree(workdir, ignore_errors=True)
    if FAILURES:
        print(f"\n{len(FAILURES)} check(s) FAILED: {FAILURES}")
        sys.exit(1)
    print("\nAll checks passed.")
    sys.exit(0)


if __name__ == "__main__":
    main()
