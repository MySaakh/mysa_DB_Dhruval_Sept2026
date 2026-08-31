#!/usr/bin/env python3
"""Smoke test for extract_headers_and_samples.py (unmasked, Excel-only).

Generates dummy fixtures, runs the extractor as a subprocess, then verifies
the export policy and every hardening fix:

  1. exit code 0 even with control characters in the data (no late crash)
  2. no live formulas anywhere in the output (injection neutralized)
  3. sample rows per sheet <= MAX_SAMPLE_ROWS
  4. values exported AS-IS: full phone numbers / emails / GSTINs present
     unmodified (no masking)
  5. cp1252 CSV included (encoding fallback), noted in the Report sheet
  6. .txt and ~$ lock files excluded
  7. headerless file flagged as header_detected=no
  8. Index sheet lists exactly the preview sheets that exist
  9. run accounting lives in the 'Report' sheet of part001 — no CSV files
     are produced

Usage: python tests/smoke_test.py
"""

import glob
import os
import shutil
import subprocess
import sys
import tempfile

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
SCRIPT = os.path.join(HERE, "..", "extract_headers_and_samples.py")

sys.path.insert(0, HERE)
import make_fixtures  # noqa: E402

FAILURES = []


def check(label, condition, detail=""):
    status = "PASS" if condition else "FAIL"
    print(f"  [{status}] {label}" + (f" — {detail}" if detail and not condition else ""))
    if not condition:
        FAILURES.append(label)


def main():
    workdir = tempfile.mkdtemp(prefix="ehs_smoke_")
    fixtures = os.path.join(workdir, "fixtures")
    output = os.path.join(workdir, "preview.xlsx")
    make_fixtures.build(fixtures)

    print("Running extractor on fixtures...")
    proc = subprocess.run(
        [sys.executable, SCRIPT, fixtures, output],
        capture_output=True, text=True,
    )
    print(proc.stdout)
    if proc.stderr:
        print(proc.stderr)

    print("Checks:")
    check("exit code 0 despite control chars in data", proc.returncode == 0,
          f"exit={proc.returncode}")

    part1 = os.path.join(workdir, "preview_part001.xlsx")
    check("part file written", os.path.exists(part1))
    check("no CSV files produced",
          not glob.glob(os.path.join(workdir, "*.csv")),
          str(glob.glob(os.path.join(workdir, "*.csv"))))
    if FAILURES:
        finish(workdir)

    wb = openpyxl.load_workbook(part1)  # keeps formulas visible if any exist
    all_cells = []
    for name in wb.sheetnames:
        for row in wb[name].iter_rows():
            for c in row:
                if c.value is not None:
                    all_cells.append((name, c))

    check("no live formulas in output",
          all(c.data_type != "f" for _, c in all_cells))

    text = " | ".join(str(c.value) for _, c in all_cells)
    check("full mobile number exported as-is", "9000000001" in text)
    check("full email exported as-is", "owner1@example.com" in text)
    check("full GSTIN exported as-is", "27ABCDE1234F1Z5" in text)
    check("pincode readable", "400001" in text)
    check("control character stripped, value kept", "hascontrol char" in text)
    check("injected formula stored as text",
          any("HYPERLINK" in str(c.value) and c.data_type != "f"
              for _, c in all_cells))
    check(".txt content not exported", "private notes" not in text)

    # per-sheet checks
    sheets = set(wb.sheetnames)
    check("Index present", "Index" in sheets)
    check("Report sheet present", "Report" in sheets)
    check("cp1252 CSV included", any("fx4_cp1252" in s for s in sheets))
    check("lock file not processed", not any("~$" in s for s in sheets))

    fx1 = next(s for s in sheets if "fx1_clean" in s)
    ws = wb[fx1]
    # meta block: 1 header + 7 rows, then 1 column-header row, then samples
    sample_rows = ws.max_row - (1 + 7 + 1)
    check("sample rows capped at 10", sample_rows == 10,
          f"got {sample_rows}")

    idx = wb["Index"]
    idx_rows = list(idx.iter_rows(min_row=2, values_only=True))
    listed = {r[0] for r in idx_rows}
    expected = sheets - {"Index", "Report"}
    check("Index lists exactly the preview sheets that exist",
          listed == expected,
          f"listed-but-missing: {listed - sheets}, unlisted: {expected - listed}")

    by_file = {}
    for r in idx_rows:
        by_file[r[1]] = r  # file_path -> row
    fx3_row = next((r for f, r in by_file.items() if "fx3_headerless" in f), None)
    check("headerless file flagged header_detected=no",
          fx3_row is not None and fx3_row[5] == "no",
          f"row={fx3_row}")
    fx2_row = next((r for f, r in by_file.items() if "fx2_banner" in f), None)
    check("banner file header found at row 3",
          fx2_row is not None and fx2_row[4] == 3, f"row={fx2_row}")

    rep = wb["Report"]
    rep_iter = rep.iter_rows(values_only=True)
    rep_header = list(next(rep_iter))
    rep_rows = [dict(zip(rep_header, r)) for r in rep_iter]
    check("Report: txt skipped with reason",
          any(r["file_path"] and "fx7" in r["file_path"]
              and r["status"] == "skipped" for r in rep_rows))
    check("Report: cp1252 encoding noted",
          any(r["file_path"] and "fx4" in r["file_path"]
              and r["detail"] and "cp1252" in r["detail"] for r in rep_rows))
    check("Report: no errors",
          all(r["status"] != "error" for r in rep_rows),
          str([r for r in rep_rows if r["status"] == "error"]))
    check("Report: every fixture accounted for",
          all(any(r["file_path"] and fx in r["file_path"] for r in rep_rows)
              for fx in ["fx1", "fx2", "fx3", "fx4", "fx5", "fx6", "fx7"]))

    finish(workdir)


def finish(workdir):
    shutil.rmtree(workdir, ignore_errors=True)
    if FAILURES:
        print(f"\n{len(FAILURES)} check(s) FAILED: {FAILURES}")
        sys.exit(1)
    print("\nAll checks passed.")
    sys.exit(0)


if __name__ == "__main__":
    main()
