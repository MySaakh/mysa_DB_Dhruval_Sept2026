#!/usr/bin/env python3
"""
extract_headers_and_samples.py  (hardened, unmasked)
----------------------------------------------------
Recursively scan a folder for Excel/CSV files, detect the header row,
and write one output sheet per source sheet to multi-sheet Excel part
files, each containing:

  - Metadata rows (file_path, file_name, sheet_name, header_row_no, ...)
  - Detected column headers
  - Up to MAX_SAMPLE_ROWS sample data rows, exported exactly as they
    appear in the source files (no masking)

EXPORT POLICY — the only data this tool is allowed to move
==========================================================
  * Reads at most MAX_ROWS_READ_PER_SHEET (25) rows from any sheet.
  * Exports all column headers of a sheet, plus at most MAX_SAMPLE_ROWS
    (10) data rows per sheet.
  * Sample values are exported AS-IS (unmasked, by the data owner's
    decision) — treat the output workbook as containing real lead data
    and share it only with people who may see it.
  * Every value written to the output is sanitized: control characters
    are stripped and formula-leading characters (= + - @) are escaped,
    so the output workbook can never contain live formulas.
  * All outputs are Excel files — no CSVs. The run accounting is a
    "Report" sheet inside the first output part file.
  * The source files are opened read-only and never modified.
  * This script performs no network activity of any kind.

These limits are enforced with hard truncation + assertions below, not
just by reader options.

Usage:
    python extract_headers_and_samples.py <folder_path> [output_xlsx] [options]

Options:
    --include-txt       also scan .txt files (excluded by default: notes
                        and readme files otherwise leak into the output)
    --overwrite         allow replacing existing output part files
    --sheets-per-file N start a new output part file every N sheets
                        (default 500; keeps workbooks openable and
                        bounds memory)
    --limit N           process only the first N source files (test runs)

Outputs (for output name "preview.xlsx"):
    preview_part001.xlsx, preview_part002.xlsx, ...
        One Index sheet per part (always the first tab, listing exactly
        the sheets that were written) + one preview sheet per source
        sheet. preview_part001.xlsx additionally carries the "Report"
        sheet: one row per source file/sheet with ok / skipped / error
        and the reason — the complete accounting of the run.
"""

import argparse
import csv
import importlib.util
import os
import re
import sys
import warnings

import pandas as pd

warnings.filterwarnings("ignore")

# ── EXPORT POLICY constants (see module docstring) ───────────────────────────
MAX_ROWS_READ_PER_SHEET = 25   # hard cap on rows read from any one sheet
MAX_SAMPLE_ROWS = 10           # data rows exported to the preview per sheet
MAX_SCAN_ROWS = 10             # rows inspected when locating the header row
MAX_SAMPLE_CELL_CHARS = 500    # sample values truncated beyond this
MAX_HEADER_CHARS = 120         # header values truncated beyond this

# ── Supported extensions ─────────────────────────────────────────────────────
EXCEL_EXTS = {
    ".xlsx",
    ".xlsm",
    ".xlsb",
    ".xltx",
    ".xltm",
    ".xls",
    ".xlam",
    ".ods",
}
CSV_EXTS = {".csv", ".tsv"}
TXT_EXTS = {".txt"}

# Extra engines needed beyond openpyxl: module to check -> pip package name
ENGINE_REQS = {
    ".xls": ("xlrd", "xlrd"),
    ".xlsb": ("pyxlsb", "pyxlsb"),
    ".ods": ("odf", "odfpy"),
}

CSV_ENCODINGS = ["utf-8-sig", "utf-8", "cp1252", "latin-1"]

# ── Output sanitization ──────────────────────────────────────────────────────
# Characters openpyxl refuses (control chars) — one of these anywhere in the
# corpus crashed the original version of this script at the final write.
try:
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE  # type: ignore
except Exception:  # pragma: no cover - fallback if openpyxl relocates it
    ILLEGAL_CHARACTERS_RE = re.compile(
        r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x84\x86-\x9f﷐-﷟￾￿]"
    )

FORMULA_LEADERS = ("=", "+", "-", "@", "\t", "\r")


def is_blank(v):
    if v is None:
        return True
    if isinstance(v, float) and v != v:  # NaN
        return True
    return str(v).strip().lower() in ("", "nan", "none")


def cell_str(v):
    return "" if is_blank(v) else str(v).strip()


def safe_cell(v, max_chars=MAX_SAMPLE_CELL_CHARS):
    """Strip illegal characters and neutralize formula interpretation.
    This is corruption/injection protection, NOT masking — the value's
    readable content is preserved."""
    s = cell_str(v)
    s = ILLEGAL_CHARACTERS_RE.sub("", s)
    if len(s) > max_chars:
        s = s[: max_chars - 1] + "…"
    if s.startswith(FORMULA_LEADERS):
        s = "'" + s  # stored as literal text, never as a formula
    return s


def df_map(df, fn):
    try:
        return df.map(fn)  # pandas >= 2.1
    except AttributeError:  # pragma: no cover - older pandas
        return df.applymap(fn)


# ── Header detection ─────────────────────────────────────────────────────────
HEADER_TOKENS = {
    "name", "company", "firm", "business", "owner", "proprietor", "contact",
    "person", "phone", "mobile", "tel", "telephone", "email", "mail",
    "address", "area", "city", "district", "state", "country", "pin",
    "pincode", "zip", "code", "no", "sr", "id", "gst", "gstin", "pan",
    "product", "products", "service", "services", "item", "items",
    "industry", "category", "type", "sector", "segment", "website", "web",
    "url", "turnover", "employees", "designation", "remarks", "status",
    "date", "location", "region", "zone", "landmark", "locality", "taluka",
    "town", "village", "dist", "branch", "office", "fax", "std", "landline",
    "whatsapp", "group", "grade", "level", "description", "dealer",
    "supplier", "exporter", "manufacturer", "keywords",
}

NUMERICISH_RE = re.compile(r"[\d\s\-\+\(\)\.,/%]*\d[\d\s\-\+\(\)\.,/%]*")


def _numericish(s):
    return bool(NUMERICISH_RE.fullmatch(s))


def _tokens(s):
    return {t for t in re.split(r"[^a-z0-9]+", s.lower()) if t}


def header_score(cells):
    vals = [cell_str(v) for v in cells]
    non_empty = [v for v in vals if v]
    if not non_empty:
        return float("-inf")
    n = len(non_empty)
    texty = sum(1 for v in non_empty if not _numericish(v) and len(v) <= 60)
    numeric = n - texty
    keyword_hits = sum(1 for v in non_empty if HEADER_TOKENS & _tokens(v))
    unique_bonus = 2.0 if len({v.lower() for v in non_empty}) == n else 0.0
    score = 2.0 * texty - 3.0 * numeric + 3.0 * keyword_hits + unique_bonus
    score += min(n, 10) * 0.25
    if n == 1:  # banner/title rows
        score -= 4.0
    return score


def find_header_row(df_raw):
    """Return (row_index, confident). Best-scoring row in the scan window;
    confident=False means the row is likely data, not a real header."""
    best_i, best_score = 0, float("-inf")
    for i in range(min(MAX_SCAN_ROWS, len(df_raw))):
        s = header_score(df_raw.iloc[i].tolist())
        if s > best_score:
            best_i, best_score = i, s
    return best_i, best_score >= 4.0


def make_column_names(header_row):
    """Build unique, non-empty column names from a header row."""
    col_names = []
    seen = {}
    for i, v in enumerate(header_row):
        s = safe_cell(v, max_chars=MAX_HEADER_CHARS)
        if not s:
            s = f"Col_{i + 1}"
        if s in seen:
            seen[s] += 1
            s = f"{s}_{seen[s]}"
        else:
            seen[s] = 0
        col_names.append(s)
    return col_names


def build_sample_table(df_raw, hrow):
    """Return sample DataFrame: detected headers + up to MAX_SAMPLE_ROWS rows,
    sanitized but otherwise exactly as found in the source."""
    col_names = make_column_names(df_raw.iloc[hrow].tolist())

    data_start = hrow + 1
    data_end = min(data_start + MAX_SAMPLE_ROWS, len(df_raw))

    if data_start >= len(df_raw):
        sample = pd.DataFrame(columns=col_names)
    else:
        sample = df_raw.iloc[data_start:data_end].copy()
        sample.columns = col_names

    sample = df_map(sample, safe_cell)

    # POLICY enforcement: never export more sample rows than allowed.
    sample = sample.iloc[:MAX_SAMPLE_ROWS]
    assert len(sample) <= MAX_SAMPLE_ROWS, "export policy violation (rows)"
    return sample


def sanitize_sheet_name(name, used_names):
    """Excel sheet names: max 31 chars, no : \\ / ? * [ ]"""
    name = re.sub(r"[:\\/?*\[\]]", "_", str(name)).strip().strip("'")
    name = name[:31] if name else "Sheet"

    base = name
    counter = 1
    while name in used_names:
        suffix = f"_{counter}"
        name = (
            base[: 31 - len(suffix)] + suffix
            if len(base) + len(suffix) > 31
            else base + suffix
        )
        counter += 1

    used_names.add(name)
    return name


# ── File discovery ───────────────────────────────────────────────────────────
def find_files(root_folder, include_txt, output_stem):
    """Yield (filepath, ext). Skips hidden dirs/files, Excel lock files, and
    this run's own output part files (in case output sits inside the scan
    folder)."""
    active_exts = EXCEL_EXTS | CSV_EXTS | (TXT_EXTS if include_txt else set())
    part_re = re.compile(re.escape(output_stem) + r"_part\d{3}\.xlsx$")
    for dirpath, dirnames, filenames in os.walk(root_folder):
        dirnames[:] = [d for d in dirnames if not d.startswith(".")]
        for fname in sorted(filenames):
            if fname.startswith(".") or fname.startswith("~$"):
                continue
            filepath = os.path.join(dirpath, fname)
            if part_re.search(os.path.abspath(filepath)):
                continue
            ext = os.path.splitext(fname)[1].lower()
            if ext in active_exts:
                yield filepath, ext
            elif ext in TXT_EXTS:
                yield filepath, "SKIP_TXT"


def missing_engine(ext):
    """Return the missing pip package name for this extension, or None."""
    if ext in ENGINE_REQS:
        module, package = ENGINE_REQS[ext]
        if importlib.util.find_spec(module) is None:
            return package
    return None


# ── Per-sheet extraction ─────────────────────────────────────────────────────
def read_excel_preview(filepath, ext):
    engine = None
    if ext == ".xls":
        engine = "xlrd"
    elif ext == ".xlsb":
        engine = "pyxlsb"
    elif ext == ".ods":
        engine = "odf"

    kwargs = dict(
        sheet_name=None, header=None, dtype=str, nrows=MAX_ROWS_READ_PER_SHEET
    )
    if engine:
        kwargs["engine"] = engine
    return pd.read_excel(filepath, **kwargs)


def extract_from_excel(filepath, ext):
    all_sheets = read_excel_preview(filepath, ext)
    for sheet_name, df in all_sheets.items():
        # POLICY enforcement: defensive truncation on top of nrows.
        df = df.iloc[:MAX_ROWS_READ_PER_SHEET]
        assert len(df) <= MAX_ROWS_READ_PER_SHEET, "export policy violation"
        if df.empty:
            continue
        hrow, confident = find_header_row(df)
        sample_df = build_sample_table(df, hrow)
        if sample_df.shape[1] > 0:
            yield sheet_name, hrow + 1, confident, sample_df, ""


def extract_from_csv(filepath, ext):
    sep = "\t" if ext == ".tsv" else ","
    try:
        with open(filepath, "r", encoding="utf-8", errors="replace") as f:
            sample_text = f.read(4096)
        dialect = csv.Sniffer().sniff(sample_text, delimiters=",\t;|")
        sep = dialect.delimiter
    except Exception:
        pass

    df, used_encoding, last_error = None, None, None
    for encoding in CSV_ENCODINGS:
        try:
            df = pd.read_csv(
                filepath,
                sep=sep,
                header=None,
                dtype=str,
                encoding=encoding,
                nrows=MAX_ROWS_READ_PER_SHEET,
                engine="python",
                on_bad_lines="skip",
            )
            used_encoding = encoding
            break
        except UnicodeDecodeError as e:
            last_error = e
        except Exception as e:
            last_error = e
            break
    if df is None:
        raise RuntimeError(f"could not read as CSV: {last_error}")

    df = df.iloc[:MAX_ROWS_READ_PER_SHEET]
    assert len(df) <= MAX_ROWS_READ_PER_SHEET, "export policy violation"
    if df.empty:
        return

    hrow, confident = find_header_row(df)
    sample_df = build_sample_table(df, hrow)
    if sample_df.shape[1] > 0:
        note = f"encoding={used_encoding}" if used_encoding != "utf-8-sig" else ""
        yield "CSV", hrow + 1, confident, sample_df, note


# ── Output writing ───────────────────────────────────────────────────────────
INDEX_COLUMNS = [
    "output_sheet", "file_path", "file_name", "source_sheet_name",
    "header_row_no", "header_detected", "total_columns", "sample_rows", "note",
]

REPORT_COLUMNS = [
    "file_path", "sheet_name", "status", "detail", "output_file",
    "output_sheet", "header_row_no", "header_detected",
    "total_columns", "sample_rows",
]

RESERVED_SHEETS = {"Index", "Report"}


def write_sheet_with_metadata(writer, sheet_name, meta, sample_df):
    """Write metadata block, then header + sample data."""
    meta_rows = [
        ["file_path", safe_cell(meta["file_path"])],
        ["file_name", safe_cell(meta["file_name"])],
        ["sheet_name", safe_cell(meta["sheet_name"])],
        ["header_row_no", meta["header_row_no"]],
        ["header_detected", "yes" if meta["header_detected"] else "NO (best guess — row may be data)"],
        ["total_columns", meta["total_columns"]],
        ["sample_rows", meta["sample_rows"]],
    ]

    meta_df = pd.DataFrame(meta_rows, columns=["Field", "Value"])
    meta_df.to_excel(writer, sheet_name=sheet_name, index=False, startrow=0)

    start_row = len(meta_rows) + 1
    sample_df.to_excel(writer, sheet_name=sheet_name, index=False, startrow=start_row)


def write_part(part_path, entries):
    """Write one output part file. The Index sheet is built from the sheets
    actually written (never from intentions) and placed first."""
    written = []
    with pd.ExcelWriter(part_path, engine="openpyxl") as writer:
        for entry in entries:
            try:
                write_sheet_with_metadata(
                    writer, entry["output_sheet"], entry["meta"], entry["sample_df"]
                )
                written.append(entry)
            except Exception as e:  # pragma: no cover - post-sanitize safety net
                entry["write_error"] = str(e)
                book = writer.book
                if entry["output_sheet"] in book.sheetnames:
                    del book[entry["output_sheet"]]

        index_rows = []
        for entry in written:
            meta = entry["meta"]
            index_rows.append({
                "output_sheet": entry["output_sheet"],
                "file_path": safe_cell(meta["file_path"]),
                "file_name": safe_cell(meta["file_name"]),
                "source_sheet_name": safe_cell(meta["sheet_name"]),
                "header_row_no": meta["header_row_no"],
                "header_detected": "yes" if meta["header_detected"] else "no",
                "total_columns": meta["total_columns"],
                "sample_rows": meta["sample_rows"],
                "note": safe_cell(meta["note"]),
            })
        pd.DataFrame(index_rows, columns=INDEX_COLUMNS).to_excel(
            writer, sheet_name="Index", index=False
        )
        try:
            book = writer.book
            book.move_sheet("Index", offset=-(len(book.sheetnames) - 1))
        except Exception:
            pass
    return written


def write_report_sheet(first_part_path, report_rows):
    """Add the run accounting as a 'Report' sheet inside the first part file
    (created standalone if no preview sheets were extracted at all)."""
    report_df = pd.DataFrame(report_rows, columns=REPORT_COLUMNS)
    if os.path.exists(first_part_path):
        with pd.ExcelWriter(first_part_path, engine="openpyxl", mode="a") as writer:
            report_df.to_excel(writer, sheet_name="Report", index=False)
            try:
                book = writer.book
                book.move_sheet("Report", offset=-(len(book.sheetnames) - 2))
            except Exception:
                pass
    else:
        with pd.ExcelWriter(first_part_path, engine="openpyxl") as writer:
            report_df.to_excel(writer, sheet_name="Report", index=False)


# ── Main ─────────────────────────────────────────────────────────────────────
def run(folder, output_xlsx, include_txt=False, overwrite=False,
        sheets_per_file=500, limit=0):
    folder = os.path.abspath(folder)
    if not os.path.isdir(folder):
        sys.exit(f"ERROR: '{folder}' is not a valid directory.")

    output_xlsx = os.path.abspath(output_xlsx)
    stem, _ = os.path.splitext(output_xlsx)

    def part_path(n):
        return f"{stem}_part{n:03d}.xlsx"

    if not overwrite and os.path.exists(part_path(1)):
        sys.exit(
            f"ERROR: output already exists ({part_path(1)}). "
            "Pass --overwrite to replace it."
        )

    missing = {}          # ext -> package, reported once
    part_entries = []
    part_no = 0
    parts_written = []
    used_sheet_names = set(RESERVED_SHEETS)
    counts = {"ok": 0, "skipped": 0, "error": 0}
    report_rows = []

    def log(rel_path, sheet, status, detail="", out_file="", out_sheet="",
            hrow="", detected="", cols="", rows=""):
        counts[status] += 1
        report_rows.append({
            "file_path": safe_cell(rel_path),
            "sheet_name": safe_cell(sheet),
            "status": status,
            "detail": safe_cell(detail),
            "output_file": out_file,
            "output_sheet": out_sheet,
            "header_row_no": hrow,
            "header_detected": detected,
            "total_columns": cols,
            "sample_rows": rows,
        })

    def flush_part():
        nonlocal part_entries, part_no
        if not part_entries:
            return
        part_no += 1
        path = part_path(part_no)
        written = write_part(path, part_entries)
        parts_written.append(path)
        for entry in part_entries:
            meta = entry["meta"]
            if entry.get("write_error"):
                log(meta["file_path"], meta["sheet_name"], "error",
                    f"write failed: {entry['write_error']}")
            else:
                log(meta["file_path"], meta["sheet_name"], "ok", meta["note"],
                    os.path.basename(path), entry["output_sheet"],
                    meta["header_row_no"],
                    "yes" if meta["header_detected"] else "no",
                    meta["total_columns"], meta["sample_rows"])
        print(f"  >> wrote {os.path.basename(path)} ({len(written)} sheets)")
        part_entries = []

    all_files = list(find_files(folder, include_txt, stem))
    if limit:
        all_files = all_files[:limit]
    total = len(all_files)
    print(f"Scanning {total} file(s) under: {folder}")
    print(f"Policy: read <= {MAX_ROWS_READ_PER_SHEET} rows/sheet, export headers "
          f"+ <= {MAX_SAMPLE_ROWS} sample rows/sheet, values exported AS-IS "
          f"(no masking)\n")

    try:
        for i, (filepath, ext) in enumerate(all_files, 1):
            rel_path = os.path.relpath(filepath, folder)
            fname = os.path.basename(filepath)
            stem_name = os.path.splitext(fname)[0]
            print(f"[{i}/{total}] {rel_path}")

            if ext == "SKIP_TXT":
                print("  [SKIP] .txt excluded by default (use --include-txt)")
                log(rel_path, "", "skipped",
                    ".txt excluded by default (use --include-txt)")
                continue

            package = missing_engine(ext)
            if package:
                if ext not in missing:
                    missing[ext] = package
                    print(f"  [WARN] {ext} files need '{package}' "
                          f"(pip install {package}) — skipping all {ext} files")
                log(rel_path, "", "skipped", f"missing engine: {package}")
                continue

            try:
                gen = (
                    extract_from_excel(filepath, ext)
                    if ext in EXCEL_EXTS
                    else extract_from_csv(filepath, ext)
                )
                got_any = False
                for sheet_name, header_row_no, confident, sample_df, note in gen:
                    got_any = True
                    out_sheet = sanitize_sheet_name(
                        f"{stem_name}_{sheet_name}", used_sheet_names
                    )
                    part_entries.append({
                        "output_sheet": out_sheet,
                        "sample_df": sample_df,
                        "meta": {
                            "file_path": rel_path,
                            "file_name": fname,
                            "sheet_name": sheet_name,
                            "header_row_no": header_row_no,
                            "header_detected": confident,
                            "total_columns": sample_df.shape[1],
                            "sample_rows": len(sample_df),
                            "note": note,
                        },
                    })
                    print(f"  [OK] '{sheet_name}' -> '{out_sheet}' | "
                          f"header row {header_row_no}"
                          f"{'' if confident else ' (LOW CONFIDENCE)'} | "
                          f"{sample_df.shape[1]} cols | "
                          f"{len(sample_df)} sample rows")
                    if len(part_entries) >= sheets_per_file:
                        flush_part()
                if not got_any:
                    log(rel_path, "", "skipped", "no readable sheets/rows")
            except Exception as e:
                print(f"  [ERROR] {rel_path}: {e}")
                log(rel_path, "", "error", str(e)[:300])
    finally:
        # Even on an interrupted run, flush collected sheets and write the
        # accounting so the outputs on disk always explain themselves.
        flush_part()
        if report_rows:
            write_report_sheet(part_path(1), report_rows)
            if part_path(1) not in parts_written:
                parts_written.insert(0, part_path(1))

    print(f"\nDone. sheets ok: {counts['ok']} | skipped: {counts['skipped']} "
          f"| errors: {counts['error']}")
    if parts_written:
        print("Output files:")
        for p in parts_written:
            print(f"  {p}")
        print(f"Run accounting: 'Report' sheet inside "
              f"{os.path.basename(part_path(1))}")
    else:
        print("No files found. Check folder path and file formats.")


def main():
    parser = argparse.ArgumentParser(
        description="Extract headers + sample rows from Excel/CSV files "
                    "(see module docstring for the export policy)."
    )
    parser.add_argument("folder", help="root folder to scan")
    parser.add_argument("output", nargs="?",
                        default="extracted_headers_and_samples.xlsx",
                        help="output name; parts written as <name>_partNNN.xlsx")
    parser.add_argument("--include-txt", action="store_true",
                        help="also scan .txt files (off by default)")
    parser.add_argument("--overwrite", action="store_true",
                        help="replace existing output files")
    parser.add_argument("--sheets-per-file", type=int, default=500,
                        help="max preview sheets per output part (default 500)")
    parser.add_argument("--limit", type=int, default=0,
                        help="process only the first N files (test runs)")
    args = parser.parse_args()

    run(args.folder, args.output,
        include_txt=args.include_txt,
        overwrite=args.overwrite,
        sheets_per_file=max(1, args.sheets_per_file),
        limit=args.limit)


if __name__ == "__main__":
    main()
